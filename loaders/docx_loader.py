# loaders/docx_loader.py
# python-docx reads .docx files paragraph by paragraph. We combine all
# paragraphs (and any tables) into a single Document per file.

from docx import Document as DocxDocument
from loaders.base import make_document


def load_docx(file_path):
    docx = DocxDocument(file_path)
    full_text = ""

    for para in docx.paragraphs:
        if para.text.strip():
            full_text += para.text + "\n"

    # Tables in Word docs aren't picked up by .paragraphs, so walk them separately.
    for table in docx.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            full_text += row_text + "\n"

    return [make_document(full_text, file_path, "DOCX")]
