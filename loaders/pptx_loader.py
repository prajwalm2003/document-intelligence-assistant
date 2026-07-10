# loaders/pptx_loader.py
# Each slide becomes its own Document, so retrieval can point back
# to a specific slide number rather than the whole deck.

from pptx import Presentation
from loaders.base import make_document


def load_pptx(file_path):
    prs = Presentation(file_path)
    docs = []

    for slide_num, slide in enumerate(prs.slides):
        slide_text = f"Slide {slide_num + 1}:\n"
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text += shape.text + "\n"

        if slide_text.strip():
            docs.append(make_document(slide_text, file_path, "PPTX", page=slide_num))

    return docs
